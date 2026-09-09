from pathlib import Path

from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter

from icloud_docs_mcp.extract import extract_file


def test_extract_txt(tmp_path: Path):
    path = tmp_path / "note.txt"
    path.write_text("hello drive", encoding="utf-8")
    result = extract_file(path)
    assert result.text == "hello drive"
    assert result.truncated is False
    assert result.content_type == "txt"


def test_extract_docx(tmp_path: Path):
    path = tmp_path / "letter.docx"
    document = Document()
    document.add_paragraph("Quote total is 1200")
    document.save(path)
    result = extract_file(path)
    assert "1200" in result.text
    assert result.content_type == "docx"


def test_extract_xlsx(tmp_path: Path):
    path = tmp_path / "sheet.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "item"
    sheet["B1"] = "qty"
    sheet["A2"] = "brick"
    sheet["B2"] = 40
    workbook.save(path)
    result = extract_file(path)
    assert "brick" in result.text
    assert result.content_type == "xlsx"


def test_extract_pdf(tmp_path: Path):
    path = tmp_path / "page.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    with path.open("wb") as handle:
        writer.write(handle)
    result = extract_file(path)
    assert result.content_type == "pdf"
    assert result.note is None


def test_extract_unsupported_pages(tmp_path: Path):
    path = tmp_path / "file.pages"
    path.write_bytes(b"not a real pages file")
    result = extract_file(path)
    assert result.text == ""
    assert result.note is not None
    assert "Pages" in result.note


def test_extract_truncates(tmp_path: Path):
    path = tmp_path / "big.txt"
    path.write_text("x" * 20_000, encoding="utf-8")
    result = extract_file(path, limit=100)
    assert result.truncated is True
    assert len(result.text) == 100
