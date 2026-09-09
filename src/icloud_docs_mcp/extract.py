"""Extract text from common document types. Never return file bytes."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

TEXT_CHAR_LIMIT = 12_000

TEXT_EXTENSIONS = {".txt", ".md", ".csv", ".tsv", ".rtf", ".log", ".json", ".xml"}
PDF_EXTENSIONS = {".pdf"}
DOCX_EXTENSIONS = {".docx"}
XLSX_EXTENSIONS = {".xlsx", ".xlsm"}
UNSUPPORTED_HINTS = {
    ".doc": "Legacy .doc is not extracted. Convert to .docx or .pdf.",
    ".ppt": "Legacy .ppt is not extracted. Convert to .pptx or .pdf.",
    ".pptx": "PowerPoint extraction is not included in v1. Export to PDF.",
    ".pages": "Apple Pages files are not extracted. Export to PDF or DOCX.",
    ".numbers": "Apple Numbers files are not extracted. Export to PDF or XLSX.",
    ".key": "Apple Keynote files are not extracted. Export to PDF.",
}


@dataclass
class ExtractResult:
    text: str
    truncated: bool
    content_type: str
    note: str | None = None


def _clip(text: str, limit: int = TEXT_CHAR_LIMIT) -> tuple[str, bool]:
    if len(text) <= limit:
        return text, False
    return text[:limit], True


def _read_text_file(path: Path) -> str:
    raw = path.read_bytes()
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _extract_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    parts: list[str] = []
    for index, page in enumerate(reader.pages, start=1):
        page_text = page.extract_text() or ""
        parts.append(f"--- page {index} ---\n{page_text}".rstrip())
        if sum(len(p) for p in parts) >= TEXT_CHAR_LIMIT:
            break
    return "\n\n".join(parts)


def _extract_docx(path: Path) -> str:
    from docx import Document

    document = Document(str(path))
    parts: list[str] = []
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            parts.append(paragraph.text)
    for table in document.tables:
        rows = []
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            rows.append(" | ".join(cells))
        if rows:
            parts.append("\n".join(rows))
    return "\n".join(parts)


def _extract_xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(str(path), read_only=True, data_only=True)
    parts: list[str] = []
    for sheet in workbook.worksheets:
        parts.append(f"--- sheet {sheet.title} ---")
        row_count = 0
        for row in sheet.iter_rows(values_only=True):
            values = ["" if cell is None else str(cell) for cell in row]
            if not any(values):
                continue
            parts.append("\t".join(values))
            row_count += 1
            if row_count >= 200:
                parts.append("[sheet truncated at 200 rows]")
                break
        if sum(len(p) for p in parts) >= TEXT_CHAR_LIMIT:
            break
    return "\n".join(parts)


def extract_file(path: Path, limit: int = TEXT_CHAR_LIMIT) -> ExtractResult:
    suffix = path.suffix.lower()
    if suffix in UNSUPPORTED_HINTS:
        return ExtractResult(
            text="",
            truncated=False,
            content_type=suffix.lstrip(".") or "binary",
            note=UNSUPPORTED_HINTS[suffix],
        )
    try:
        if suffix in TEXT_EXTENSIONS or suffix == "":
            text = _read_text_file(path)
            content_type = suffix.lstrip(".") or "text"
        elif suffix in PDF_EXTENSIONS:
            text = _extract_pdf(path)
            content_type = "pdf"
        elif suffix in DOCX_EXTENSIONS:
            text = _extract_docx(path)
            content_type = "docx"
        elif suffix in XLSX_EXTENSIONS:
            text = _extract_xlsx(path)
            content_type = "xlsx"
        else:
            return ExtractResult(
                text="",
                truncated=False,
                content_type=suffix.lstrip(".") or "binary",
                note=f"No text extractor for {suffix or 'this file type'}.",
            )
    except Exception as exc:  # noqa: BLE001 — surface extractor failures to the tool
        return ExtractResult(
            text="",
            truncated=False,
            content_type=suffix.lstrip(".") or "unknown",
            note=f"Extraction failed: {type(exc).__name__}: {exc}",
        )

    clipped, truncated = _clip(text, limit)
    return ExtractResult(text=clipped, truncated=truncated, content_type=content_type)
